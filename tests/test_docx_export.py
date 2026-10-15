"""DOCX export: single column, real headings, every bullet in order."""

from docx import Document

from strelitzia.scoring.docx_export import to_docx
from strelitzia.scoring.resume_model import TailoredResume
from strelitzia.scoring.template import fixed_roles_from_resume, parse_base_resume
from workbench_helpers import BASE_RESUME, FIXED, PROFILE, good_result, library


def _model():
    lib = library()
    return TailoredResume.from_llm_json(good_result()["resume"], lib,
                                        fixed_roles=fixed_roles_from_resume(parse_base_resume(BASE_RESUME, lib)))


def test_docx_headings_and_bullets_in_order(tmp_path):
    model = _model()
    path = to_docx(model, tmp_path / "r.docx", PROFILE, FIXED)
    doc = Document(str(path))
    paras = [(p.style.name, p.text) for p in doc.paragraphs]
    headings = [(s, t) for s, t in paras if s.startswith(("Heading", "Title"))]
    assert headings == [
        ("Title", "Alex Tester"),
        ("Heading 1", "Work Experience"),
        ("Heading 2", "Data Engineer at Acme Corp · Jan 2024 – Present"),
        ("Heading 2", "Data Analyst Intern at Globex · Jun 2022 – Aug 2022"),
        ("Heading 1", "Education"),
        ("Heading 2", "State University, B.S. Statistics · May 2022"),
        ("Heading 1", "Skills"),
    ]
    bullets = [t for s, t in paras if s == "List Bullet"]
    expected = [b.text for r in model.roles for b in r.bullets] + FIXED["education"][0]["bullets"]
    assert bullets == expected
    texts = [t for _, t in paras]
    assert model.headline in texts and "Springfield, IL | alex@example.com" in texts
    assert "Languages: Python, SQL" in texts and "Tools: Airflow, Spark, dbt, Superset" in texts


def test_docx_has_no_tables_or_header_footer_text(tmp_path):
    doc = Document(str(to_docx(_model(), tmp_path / "r.docx", PROFILE, FIXED)))
    assert doc.tables == []
    for section in doc.sections:
        assert all(not p.text.strip() for p in section.header.paragraphs)
        assert all(not p.text.strip() for p in section.footer.paragraphs)
    assert not doc.inline_shapes
