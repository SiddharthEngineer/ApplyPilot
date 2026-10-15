"""Word export of a tailored resume, for ATSs that parse DOCX better than PDF.

Single column with real heading styles: no tables, text boxes, images, or header/footer content. The header
(name, contact) and education come from the same fixed blocks as the PDF (template.load_fixed).
"""

from __future__ import annotations

from pathlib import Path

from strelitzia.scoring.resume_model import TailoredResume
from strelitzia.scoring.template import build_contact, load_fixed


def to_docx(resume: TailoredResume, path: Path | str, profile: dict | None = None, fixed: dict | None = None) -> Path:
    """Write the resume as a .docx at `path` and return the path."""
    from docx import Document
    from docx.shared import Inches, Pt

    profile = profile or {}
    fixed = fixed if fixed is not None else load_fixed()
    header = fixed.get("header") or {}
    name = header.get("name") or profile.get("personal", {}).get("full_name", "")
    contact = header.get("contact") or build_contact(profile)

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.5)
        section.left_margin = section.right_margin = Inches(0.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    doc.add_heading(name, level=0)
    if resume.headline:
        doc.add_paragraph(resume.headline)
    if contact:
        doc.add_paragraph(" | ".join(c["text"] for c in contact))

    doc.add_heading("Work Experience", level=1)
    for role in resume.roles:
        head = f"{role.title} at {role.company}" if role.company else role.title
        doc.add_heading(f"{head} · {role.dates}" if role.dates else head, level=2)
        if role.tagline:
            doc.add_paragraph(role.tagline).runs[0].italic = True
        for b in role.bullets:
            doc.add_paragraph(b.text, style="List Bullet")

    if fixed.get("education"):
        doc.add_heading("Education", level=1)
        for e in fixed["education"]:
            doc.add_heading(f"{e['heading']} · {e['dates']}" if e["dates"] else e["heading"], level=2)
            for b in e["bullets"]:
                doc.add_paragraph(b, style="List Bullet")

    if resume.skills:
        doc.add_heading("Skills", level=1)
        for cat, items in resume.skills.items():
            p = doc.add_paragraph()
            p.add_run(f"{cat}: ").bold = True
            p.add_run(items)

    path = Path(path)
    doc.save(str(path))
    return path
