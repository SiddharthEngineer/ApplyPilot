"""Resume upload on the Profile page: the text of a PDF or DOCX, and Claude's draft of the profile fields it states.

The draft only holds what the resume says (no guesses); the web layer merges it into the empty fields for the user
to review. Never log the resume text.
"""

from __future__ import annotations

import io
import logging

from strelitzia.llm import acting_for, check_budget, get_client, json_object, nullable
from strelitzia.llm_json import extract_json
from strelitzia.profile import PDF_MIME, PROFILE_EDUCATION_LEVELS

log = logging.getLogger(__name__)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_RESUME_CHARS = 30_000  # what reaches the prompt; a resume is a few thousand
_MAX_OUTPUT_TOKENS = 2000

# The profile fields a resume can state (UserProfile's names).
TEXT_FIELDS = ("first_name", "last_name", "phone", "email", "linkedin_url", "github_url", "location",
               "current_job_title")

PARSE_SCHEMA = json_object({
    **{name: nullable({"type": "string"}) for name in TEXT_FIELDS},
    "years_experience": nullable({"type": "number"}),
    "education_level": nullable({"type": "string", "enum": list(PROFILE_EDUCATION_LEVELS)}),
    "skills": {"type": "array", "items": {"type": "string"}},
})

SYSTEM_PROMPT = f"""You read a resume and return the candidate's profile fields as JSON.

CORE RULE: extract, never infer. Use null (or [] for skills) for anything the resume doesn't state.
Copy names, contact details and titles as written.

- first_name / last_name: the candidate's name from the header.
- phone, email: as written. linkedin_url / github_url: the profile URLs, with https:// added when the resume
  leaves it out ("linkedin.com/in/x" -> "https://linkedin.com/in/x"). null when absent.
- location: where the candidate lives, from the header ("Austin, TX"). Not an employer's office.
- current_job_title: the title of the most recent (or current) position.
- years_experience: total years of professional experience, as a number. Use a figure the resume states
  ("8+ years" -> 8); otherwise add up the dated positions (internships and education don't count). null when
  there are no dates.
- education_level: the highest completed degree, one of {' | '.join(PROFILE_EDUCATION_LEVELS)}.
  null when the resume has no education section.
- skills: tools, languages, frameworks and methods listed in the resume ("Python", "SQL", "Kubernetes").
  One per item, no duplicates, at most 40.
"""


def extract_text(name: str, mime: str | None, data: bytes) -> str:
    """The text of a PDF or DOCX resume (picked by extension or MIME type).

    Raises:
        ValueError: Not a PDF or DOCX, a file that can't be read, or one with no text (a scanned PDF).
    """
    lower = (name or "").lower()
    if lower.endswith(".pdf") or mime == PDF_MIME:
        text = _pdf_text(data)
    elif lower.endswith(".docx") or mime == DOCX_MIME:
        text = _docx_text(data)
    else:
        raise ValueError("Upload a PDF or DOCX resume")
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ValueError("No text found in the file (a scanned PDF needs a text layer)")
    return text


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except (PyPdfError, ValueError, KeyError, TypeError, OSError) as e:
        raise ValueError(f"Couldn't read the PDF ({type(e).__name__})") from e


def _docx_text(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # python-docx raises zipfile, lxml and its own errors for a bad file
        raise ValueError(f"Couldn't read the DOCX ({type(e).__name__})") from e
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def parse_resume(text: str, user_id: str | None, conn=None) -> dict:
    """The profile fields `text` states, as UserProfile names (fields the resume doesn't state are absent).

    Billed to `user_id` and held to their cap: raises LLMStopRun("budget") over it, LLMStopRun("auth") without a key.
    Raises LLMResponseError or ValueError for an unusable answer.
    """
    check_budget(conn=conn, user_id=user_id)
    messages = [{"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": f"<resume>\n{text[:MAX_RESUME_CHARS]}\n</resume>"}]
    with acting_for(user_id):
        reply = get_client("profile_parse").chat(messages, max_tokens=_MAX_OUTPUT_TOKENS, schema=PARSE_SCHEMA)
    data = extract_json(reply)
    if not isinstance(data, dict):
        raise ValueError("The resume parse answer wasn't a JSON object")
    return clean_draft(data)


def clean_draft(data: dict) -> dict:
    """Keep the stated, well-typed fields: non-empty strings, a number of years, a known education level, skills."""
    out: dict = {}
    for name in TEXT_FIELDS:
        value = data.get(name)
        if isinstance(value, str) and value.strip():
            out[name] = value.strip()
    years = data.get("years_experience")
    if isinstance(years, (int, float)) and not isinstance(years, bool) and 0 <= years <= 70:
        out["years_experience"] = float(years)
    if data.get("education_level") in PROFILE_EDUCATION_LEVELS:
        out["education_level"] = data["education_level"]
    skills: list[str] = []
    for s in data.get("skills") if isinstance(data.get("skills"), list) else []:
        s = str(s).strip()
        if s and s.lower() not in {x.lower() for x in skills}:
            skills.append(s)
    if skills:
        out["skills"] = skills
    return out
